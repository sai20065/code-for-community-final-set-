"""Generates the Prajadhwani concept document (SFIC 2026) as a Word file.

    python docs/concept/infographics/build_infographics.py   # figures
    python docs/concept/frame_screens.py                     # app screenshots
    python docs/concept/build_docx.py                        # this document

Team details live in TEAM below — replace the placeholders and re-run.
"""
import datetime
import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = pathlib.Path(__file__).parent
FIG = HERE / "assets" / "figures"
SCR = HERE / "assets" / "screens_framed"
LOGO = HERE.parent.parent / "assets" / "icon" / "logo.png"
OUT = HERE / "Prajadhwani_Concept_Document_SFIC2026.docx"

INDIGO = RGBColor(0x2E, 0x1F, 0x8F)
INDIGO_D = RGBColor(0x1C, 0x12, 0x59)
SAFFRON_D = RGBColor(0xC9, 0x79, 0x14)
TEAL = RGBColor(0x0B, 0x8A, 0x6C)
VERM = RGBColor(0xE0, 0x38, 0x4A)
GREY = RGBColor(0x57, 0x54, 0x7A)

TEAM = {
    "Team name": "[Team name]",
    "Team members": "[Member 1 (Lead)], [Member 2], [Member 3], [Member 4]",
    "Institution": "[Institution name, City]",
    "SFIC participation level": "[Open — UG/PG/Research]  /  [Technical — ITI/Polytechnic]",
    "Institution coordinator": "[Coordinator name, designation]",
    "Contact": "[email@institution.edu.in]  ·  [+91 XXXXX XXXXX]",
}

TODAY = datetime.date(2026, 10, 22)

doc = Document()
_fig_no = [0]
_tab_no = [0]


# ----------------------------------------------------------------- helpers
def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def cell_margins(table, top=80, bottom=80, left=110, right=110):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tbl_pr.append(mar)


def borders(table, color="C9C5E3", size=4):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:color"), color)
        b.append(e)
    tbl_pr.append(b)


def no_borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tbl_pr.append(b)


def field(paragraph, instr):
    """Inserts a Word field (PAGE, TOC …) into a paragraph."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
        if kind == "separate":
            t = OxmlElement("w:t")
            t.text = "Right-click → Update field to build the table of contents." if "TOC" in instr else "1"
            run._r.append(t)
    return run


def para(text="", bold_lead=None, size=None, color=None, align=None, space_after=6, italic=False, style=None):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        if color:
            r.font.color.rgb = color
    r = p.add_run(text)
    r.italic = italic
    if size:
        for run in p.runs:
            run.font.size = Pt(size)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(2)


def h(text, level=1):
    return doc.add_heading(text, level=level)


def caption(prefix, text):
    if prefix == "Figure":
        _fig_no[0] += 1
        n = _fig_no[0]
    else:
        _tab_no[0] += 1
        n = _tab_no[0]
    p = doc.add_paragraph(style="Caption")
    r = p.add_run(f"{prefix} {n}: ")
    r.bold = True
    p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def figure(name, text, width=16.5):
    doc.add_picture(str(FIG / f"{name}.png"), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    caption("Figure", text)


def table(header, rows, widths=None, head_fill="2E1F8F", zebra=True, font=9.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    cell_margins(t)
    for i, htxt in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(htxt)
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, head_fill)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            if isinstance(val, tuple):
                r = p.add_run(val[0])
                r.bold = True
                r.font.size = Pt(font)
                if len(val) > 1:
                    r2 = p.add_run(val[1])
                    r2.font.size = Pt(font)
            else:
                r = p.add_run(str(val))
                r.font.size = Pt(font)
            if zebra and ri % 2 == 1:
                shade(cells[i], "F3F2FA")
    # Repeat header row on page breaks.
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    tr_pr.append(th)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(title, body_lines, fill="EEF0FB", edge="2E1F8F"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, color=edge, size=8)
    cell_margins(t, 140, 140, 200, 200)
    c = t.rows[0].cells[0]
    shade(c, fill)
    c.text = ""
    p = c.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = INDIGO_D
    for line in body_lines:
        q = c.add_paragraph()
        q.paragraph_format.space_after = Pt(3)
        if isinstance(line, tuple):
            rr = q.add_run(line[0])
            rr.bold = True
            rr.font.size = Pt(10)
            rr2 = q.add_run(line[1])
            rr2.font.size = Pt(10)
        else:
            rr = q.add_run(line)
            rr.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def screens_grid(items, cols=3, width=4.6):
    """items = [(screen_name, caption_text)] laid out in a borderless grid."""
    rows = (len(items) + cols - 1) // cols
    t = doc.add_table(rows=rows * 2, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t)
    for i, (name, cap) in enumerate(items):
        r, c = divmod(i, cols)
        pic_cell = t.rows[r * 2].cells[c]
        pic_cell.text = ""
        pp = pic_cell.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.add_run().add_picture(str(SCR / f"{name}.png"), width=Cm(width))
        cap_cell = t.rows[r * 2 + 1].cells[c]
        cap_cell.text = ""
        cp = cap_cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _fig_no[0] += 1
        rr = cp.add_run(f"Figure {_fig_no[0]}: ")
        rr.bold = True
        rr.font.size = Pt(8.5)
        r2 = cp.add_run(cap)
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = GREY
        cp.paragraph_format.space_after = Pt(10)
    for row in t.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------- styles & page setup
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
sec.different_first_page_header_footer = True

st = doc.styles
st["Normal"].font.name = "Calibri"
st["Normal"].font.size = Pt(11)
st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st["Normal"].paragraph_format.line_spacing = 1.12
for name, size, color in (("Heading 1", 17, INDIGO_D), ("Heading 2", 13.5, INDIGO), ("Heading 3", 11.5, SAFFRON_D)):
    s = st[name]
    s.font.name = "Cambria"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = color
    s.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True
st["Caption"].font.size = Pt(9)
st["Caption"].font.italic = False
st["Caption"].font.color.rgb = GREY

# Header / footer (from page 2).
hp = sec.header.paragraphs[0]
hp.text = ""
r = hp.add_run("Prajadhwani — Concept Document  |  Seva First Innovation Challenge 2026")
r.font.size = Pt(8.5)
r.font.color.rgb = GREY
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("For evaluation  ·  Page ")
r.font.size = Pt(8.5)
r.font.color.rgb = GREY
field(fp, "PAGE").font.size = Pt(8.5)

# ----------------------------------------------------------------- cover
band = doc.add_table(rows=1, cols=3)
no_borders(band)
for c, fill in zip(band.rows[0].cells, ("FF9933", "FFFFFF", "138808")):
    shade(c, fill)
    c.text = ""
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    c.paragraphs[0].add_run(" ").font.size = Pt(4)
doc.add_paragraph().paragraph_format.space_after = Pt(24)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(LOGO), width=Cm(4.2))

for text, size, color, bold in (
    ("PRAJADHWANI", 34, INDIGO_D, True),
    ("Voice of the People", 14, SAFFRON_D, False),
    ("AI-Powered Citizen Voice & Grievance Routing Platform", 17, INDIGO, True),
    ("with a fine-tuned Gemma 4 Responsible-Officer Router", 12.5, GREY, False),
):
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after = Pt(4)
    rr = q.add_run(text)
    rr.font.size = Pt(size)
    rr.font.color.rgb = color
    rr.bold = bold
    rr.font.name = "Cambria"

doc.add_paragraph().paragraph_format.space_after = Pt(10)
q = doc.add_paragraph()
q.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = q.add_run("CONCEPT DOCUMENT")
rr.bold = True
rr.font.size = Pt(13)
rr.font.color.rgb = INDIGO_D
q = doc.add_paragraph()
q.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = q.add_run("Submitted to the Seva First Innovation Challenge 2026 · Ministry of Education, Government of India")
rr.font.size = Pt(10.5)
rr.font.color.rgb = GREY

cover = doc.add_table(rows=0, cols=2)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
borders(cover)
cell_margins(cover)
for k, v in [("Theme", "Swachh & Sustainable Bharat (Clean and Sustainable Bharat)"),
             ("Innovation track", "Track B — Business Model & Governance (with technology core)"),
             ("Technology Readiness Level", "TRL 5 — validated in a relevant environment"),
             *TEAM.items(),
             ("Date", TODAY.strftime("%d %B %Y"))]:
    cells = cover.add_row().cells
    cells[0].text, cells[1].text = "", ""
    a = cells[0].paragraphs[0].add_run(k)
    a.bold = True
    a.font.size = Pt(10)
    a.font.color.rgb = INDIGO_D
    shade(cells[0], "EEF0FB")
    b = cells[1].paragraphs[0].add_run(v)
    b.font.size = Pt(10)
    cells[0].width, cells[1].width = Cm(5.2), Cm(11.2)
page_break()

# ----------------------------------------------------------------- document control + TOC
h("Document Control", 1)
table(["Item", "Details"], [
    ("Document title", "Prajadhwani — Concept Document"),
    ("Version", "1.0"),
    ("Date of issue", TODAY.strftime("%d %B %Y")),
    ("Classification", "For evaluation — Seva First Innovation Challenge 2026"),
    ("Prepared by", TEAM["Team name"] + ", " + TEAM["Institution"]),
    ("Status of components", "Implemented and running: Sections 3–5. Proposed (Phase 2): Section 6 (Gemma 4 router)."),
], widths=[5, 11.4])
para("All statistics are cited in Annexure C. Figures for the prototype (Section 4.4) are read live from the project's "
     "Cloud Firestore database and include seeded test data; they demonstrate that the pipeline works, not real-world adoption.",
     size=9.5, italic=True)
h("Contents", 1)
field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
page_break()

# ----------------------------------------------------------------- 1 executive summary
h("1. Executive Summary", 1)
para("Every day, citizens of Karnataka encounter civic problems — a broken streetlight, an open drain, a pothole that injures "
     "two-wheeler riders. Bengaluru alone logged more than 1.27 lakh civic grievances in the first five and a half months of 2025, "
     "over 750 every day. Yet the citizen rarely knows which of a dozen agencies is responsible, complaint channels are fragmented, "
     "and tickets are sometimes closed without the problem being fixed.")
para("Prajadhwani (“voice of the people”) is a working mobile platform that lets any citizen report a problem in 20 seconds — by "
     "voice, text or photo, in 10 Indian languages, without Aadhaar, phone number or email — and gives them an instant receipt. "
     "Artificial intelligence then understands the report, places it precisely on the map of wards, assembly seats, taluks and "
     "districts, groups it with similar reports into hotspots, and presents it on the dashboard of the official responsible for that "
     "area, together with AI-generated briefings and downloadable reports. The citizen tracks every status change live.")
callout("At a glance", [
    ("Problem: ", "fragmented, slow and opaque civic grievance redress; citizens cannot tell who is accountable."),
    ("Solution: ", "a multilingual citizen app + AI triage + jurisdiction-aware dashboards for MPs, MLAs and officers."),
    ("Core innovation (proposed): ", "a fine-tuned, open-weight Gemma 4 model that names the responsible office and officer for "
                                     "every report, with an automatic District Collector fallback."),
    ("Status: ", "TRL 5 — live cloud backend in India (Mumbai region), Android app distributed to testers, AI pipeline verified end-to-end."),
    ("Business model: ", "free for citizens; sustained through per-ward / per-district licences for government, plus public-programme funding."),
    ("Ask: ", "pilot support in one GBA ward (TRL 6), mentorship, and cloud/GPU credits to fine-tune the Gemma 4 router."),
])
figure("f04_how_it_works", "How Prajadhwani works — six steps from a citizen's report to official action")

# ----------------------------------------------------------------- 2 problem
page_break()
h("2. Problem Statement and Demographic Context", 1)
h("2.1 Scale of the challenge", 2)
para("Karnataka had 6.11 crore residents at the 2011 Census and is projected to reach about 6.9 crore in 2026. Bengaluru has grown "
     "into a metropolis of roughly 1.16 crore people within the new Greater Bengaluru Authority (GBA) area. Civic administration "
     "spans several overlapping layers: 31 districts, 224 assembly constituencies, 28 Lok Sabha constituencies, and — in Bengaluru — "
     "five new city corporations with 369 wards that replaced the 198-ward BBMP in 2025.")
figure("f01_demographics", "Demographic and administrative context of Karnataka and Bengaluru")
h("2.2 The grievance-redress gap", 2)
para("High volume is only part of the problem. Responsibility is split between municipal bodies and parastatals (electricity, water, "
     "public works, health, education), so the same pothole may belong to the city corporation or to the Public Works Department "
     "depending on the road. Citizens must know the right portal and the right office; many give up. Even national systems that "
     "perform well on averages carry large pending volumes.")
figure("f02_grievance_gap", "Grievance volumes in Bengaluru and the national CPGRAMS system")
table(["Gap", "What citizens experience", "Consequence for government"], [
    ("Who is responsible?", "Does not know whether BESCOM, BWSSB, GBA or PWD owns the issue", "Mis-routed tickets, transfers, delay"),
    ("Language & literacy", "Forms are text-heavy and mostly in English/Kannada; ~1 in 4 adults (2011) not literate", "Under-reporting from the most affected groups"),
    ("No feedback loop", "No receipt, no visible status, tickets closed without repair", "Loss of trust; duplicate complaints"),
    ("No aggregate view", "Each complaint is treated in isolation", "Officials and legislators cannot see hotspots or plan budgets"),
    ("Fragmented boundaries", "New GBA wards vs old BBMP wards vs assembly seats", "Data cannot be compared across jurisdictions"),
], widths=[3.4, 6.8, 6.2])
caption("Table", "Root causes of the grievance-redress gap")
figure("f03_gba_wards", "Ward distribution across the five Greater Bengaluru corporations", width=14.5)

# ----------------------------------------------------------------- 3 solution
page_break()
h("3. Proposed Solution: Prajadhwani", 1)
para("Prajadhwani is a civic-intelligence platform with three connected faces — one for citizens, one for government officials "
     "and elected representatives, and one public, anonymised view for everyone. It is built to be inclusive by default and "
     "accountable by design.")
h("3.1 Users and what each gets", 2)
table(["User", "Access", "Key capabilities"], [
    ("Citizen", "Username + password; no Aadhaar, phone or email",
     "Report by voice / text / photo in 10 languages; instant ticket receipt; live status tracker; civic profile with levels and badges; Google Maps directions to the reported spot"),
    ("MP / MLA", "Verified government sign-in", "Constituency dashboard, hotspot map, AI briefing on top problems, downloadable PDF report, status updates"),
    ("Ward officer / Tahsildar / District officer", "Self-registration, approved by Prajadhwani Admin",
     "Queue of tickets inside their own jurisdiction only; bulk status updates; performance statistics"),
    ("Prajadhwani Admin", "Restricted admin console", "Approves official registrations, oversees the platform"),
    ("Public / media / researchers", "No sign-in", "Anonymised public dashboard and maps — no names, addresses or exact locations"),
], widths=[3.4, 4.4, 8.6])
caption("Table", "User groups and capabilities")
h("3.2 What makes Prajadhwani different", 2)
table(["Capability", "Conventional grievance portals", "Prajadhwani"], [
    ("Input", "Typed forms, category menus", "Voice, text and photo on one screen; AI fills the category"),
    ("Languages", "One or two", "10 Indian languages with automatic translation"),
    ("Identity", "Phone / email / ID often required", "Username only — lowers barriers and protects privacy"),
    ("Location", "Citizen selects ward or office", "GPS automatically resolved to ward, assembly seat, taluk, district and Lok Sabha seat"),
    ("Routing", "Manual, by category", "AI department routing today; fine-tuned Gemma 4 officer routing proposed"),
    ("Aggregation", "Individual tickets", "AI hotspots, root-cause analysis and solution cards"),
    ("Transparency", "Limited", "Live citizen tracker + public anonymised dashboard"),
    ("Representatives", "Not connected", "MP and MLA dashboards with AI briefings and reports"),
], widths=[3.2, 5.6, 7.6])
caption("Table", "Comparison with conventional grievance systems")

# ----------------------------------------------------------------- 4 app walkthrough
page_break()
h("4. The Working Prototype — App Walkthrough", 1)
para("The screens below are real screenshots captured from the current Android build running against the live cloud backend "
     "on 22 October 2026, using a temporary demonstration account (since deleted).")
h("4.1 Inclusive onboarding", 2)
screens_grid([
    ("01_welcome", "Welcome screen with separate entry for citizens and for MPs, MLAs & officers"),
    ("02_signup", "Sign-up with only a username and password; a security question replaces phone/email recovery"),
    ("03_language", "Choice of 10 Indian languages; the whole app switches instantly"),
    ("04_location", "Home location on a free OpenStreetMap-based map (no paid licence)"),
    ("05_profile_reveal", "First-time civic profile: ward, constituency, MP and three first steps"),
    ("06b_home_active", "Citizen home: constituency, civic-activity streak and one-tap reporting"),
])
h("4.2 Reporting a problem in one screen", 2)
para("The redesigned report screen brings voice, text and photo together. The citizen can type or tap the microphone (speech is "
     "converted to editable text), attach a photo, optionally pick a category, and confirm where the problem is. A ticket number is "
     "issued the instant it is saved. On the confirmation screen an animated officer salutes and stamps the ticket, reassuring the "
     "citizen that the problem has been registered.")
screens_grid([
    ("07_compose_top", "Unified report screen: text and photo added; microphone for voice"),
    ("08_compose_bottom", "Category and location (current location, home, or drop a pin)"),
    ("10_my_reports", "My Tickets: every report with its receipt ID"),
    ("11_report_detail", "Live status tracker Filed → Acknowledged → In Progress → Resolved, with Google Maps directions"),
    ("13_public_dashboard", "Public dashboard: anonymised totals, resolution rate and top hotspots"),
    ("14_districts", "Explore Karnataka: all districts and taluks on one map"),
])
h("4.3 Every jurisdiction mapped", 2)
screens_grid([
    ("15_bengaluru_wards", "Bengaluru: both the 198 BBMP wards and the 369 GBA wards"),
    ("16_bbmp_map", "BBMP wards grouped by assembly constituency"),
], cols=3)
para("Officials sign in to a separate government dashboard scoped to their own jurisdiction, with a hotspot map, an AI briefing "
     "card summarising the most pressing problems, ticket management with bulk status updates, and a downloadable PDF report for "
     "their constituency or area.", size=10.5)
h("4.4 Prototype data snapshot", 2)
figure("f10_dataset", "Tickets in the prototype database by theme and by status (includes seeded test data)")

# ----------------------------------------------------------------- 5 technology
page_break()
h("5. Technology Architecture", 1)
para("Prajadhwani is serverless and hosted in Google Cloud's Mumbai region (asia-south1), so citizen data stays in India. "
     "The mobile app is built in Flutter; Firebase provides authentication, the Cloud Firestore database and file storage; "
     "Cloud Functions run the AI pipeline; and Gemini 2.5 Flash on Vertex AI provides language understanding.")
figure("f05_architecture", "System architecture")
h("5.1 The AI pipeline for every report", 2)
table(["Stage", "What happens", "Technology"], [
    ("1. Receipt", "Ticket ID generated in a database transaction before any AI runs, so the citizen never loses a report", "Firestore transaction"),
    ("2. Understand", "Voice transcribed (on device first, server fallback); text translated to English; photo captioned", "Gemini 2.5 Flash, Cloud Translation"),
    ("3. Classify", "Theme (if not chosen), priority, demand, demographic and infrastructure-gap scores, and a privacy-safe one-line public summary", "Gemini 2.5 Flash"),
    ("4. Locate", "GPS point resolved to Lok Sabha seat, assembly seat, district, taluk, BBMP ward and GBA ward", "Turf point-in-polygon over official boundary data"),
    ("5. Group", "Ticket joins or starts an issue cluster (same area and theme); aggregates updated", "Cloud Functions"),
    ("6. Publish", "Anonymised copy with snapped location written for the public dashboard", "Allow-listed projection"),
    ("7. Analyse hotspots", "At 5 / 15 / 40 reports or priority ≥ 70, four AI agents produce a solution card", "Multi-agent chain (Figure below)"),
], widths=[2.8, 9.2, 4.4])
caption("Table", "Processing stages applied to every report")
figure("f06_agent_chain", "Multi-agent AI pipeline for hotspots")
h("5.2 Security and privacy by design", 2)
bullets([
    ("Data minimisation: ", "no Aadhaar, phone number or email is collected from citizens."),
    ("Role-scoped access: ", "database security rules let an official read only tickets inside their own jurisdiction; officials "
                             "cannot change their own jurisdiction or role."),
    ("Public data isolation: ", "the public dashboard reads only a separate, allow-listed projection with approximate locations — "
                                "the raw tickets are never publicly readable."),
    ("Data residency: ", "all services run in the asia-south1 (Mumbai) region."),
    ("Verified officials: ", "MLAs and officers self-register and are approved by an administrator before gaining access."),
])

# ----------------------------------------------------------------- 6 gemma router
page_break()
h("6. Fine-tuned Gemma 4 Responsible-Officer Router (Proposed Core Innovation)", 1)
callout("Status", ["This component is proposed for Phase 2 (months 2–5 of the roadmap). Today, department routing is performed "
                   "per hotspot by a general-purpose model choosing from a fixed 16-department directory. The router below extends "
                   "this to every individual report and to a named officer."], fill="FFF4E2", edge="C97914")
h("6.1 Why a specialised model", 2)
para("Deciding who must act on a civic problem requires local institutional knowledge: a pothole on a GBA road belongs to the "
     "ward engineer, on a state highway to the Public Works Department; a dark street may be a BESCOM supply fault or a municipal "
     "streetlight contract. General-purpose language models do not reliably know Karnataka's allocation of business, and sending "
     "every report to a large proprietary model is costly and raises data-sovereignty concerns.")
para("Gemma 4 is Google's family of open-weight models, released in 2026 in sizes from E2B and E4B (built for on-device and edge "
     "use) to a 26B mixture-of-experts model with about 3.8B active parameters and a 31B dense model. Open weights mean the "
     "State can host, audit and own the fine-tuned model.")
figure("f07_gemma_router", "Proposed inference flow of the Responsible-Officer Router")
h("6.2 Model design", 2)
table(["Aspect", "Design decision"], [
    ("Base model", "Gemma 4 E4B (≈4B effective parameters) — low serving cost; 26B A4B MoE as the high-accuracy tier if E4B misses targets"),
    ("Task", "Structured classification: {departmentId, officerRole, jurisdictionLevel, confidence, reason}"),
    ("Input", "PII-free public summary + English translation, photo caption, theme, priority, area context (urban ward or rural taluk, district)"),
    ("Constrained output", "Department and role must come from the official enumerated lists; any other output is treated as a failure → fallback"),
    ("Officer resolution", "Deterministic lookup of the approved officer account covering the ticket's ward → taluk → district; the model never invents a person"),
    ("Fallback", "Confidence < 0.5 or no matching account → assigned to the post “District Collector, <district>” for re-assignment"),
    ("Notification", "Email and push notification to the assigned officer; “Assigned to me” queue; citizen sees the responsible designation"),
    ("Hosting", "Vertex AI / Cloud Run with GPU in asia-south1; batch or near-real-time"),
], widths=[3.6, 12.8])
caption("Table", "Design of the Gemma 4 router")
h("6.3 Training data and fine-tuning", 2)
figure("f08_training", "Fine-tuning and continuous-learning pipeline")
table(["Data source", "Volume (target)", "Purpose"], [
    ("Department directory & allocation-of-business rules", "16+ departments, all officer roles", "Ground truth for labels"),
    ("Synthetic multilingual reports", "≈20,000 examples across 10 languages and 8 themes", "Coverage of rare categories and languages"),
    ("Anonymised real tickets from the pilot", "≈5,000 examples", "Real-world phrasing and noise"),
    ("Officer re-assignments", "Continuous", "Corrections become new labels each month"),
], widths=[6.4, 5.0, 5.0])
caption("Table", "Training data plan")
para("Fine-tuning uses parameter-efficient LoRA / QLoRA supervised fine-tuning (for example rank 16, 3 epochs) on Vertex AI or "
     "NVIDIA NeMo, which fits on a single data-centre GPU. The model first runs in shadow mode beside the current system; it goes "
     "live only after meeting the targets below.")
h("6.4 Evaluation targets", 2)
table(["Metric", "Target before go-live"], [
    ("Top-1 department accuracy (held-out test set)", "≥ 90%"),
    ("Officer-role accuracy", "≥ 85%"),
    ("Share of tickets falling back to the District Collector", "≤ 15%"),
    ("Re-assignment rate after go-live", "≤ 10%"),
    ("Accuracy gap between best and worst language", "≤ 5 percentage points"),
    ("Median time from report to officer acknowledgement", "< 48 hours in the pilot ward"),
], widths=[10.4, 6.0])
caption("Table", "Acceptance criteria for the router")
h("6.5 Responsible AI safeguards", 2)
bullets([
    ("Human in the loop: ", "officers and District Collectors can re-assign; the system learns from every correction."),
    ("Explainability: ", "each assignment stores the model's short reason and confidence, visible to officials."),
    ("Auditability: ", "an immutable assignment log supports grievance appeals and RTI queries."),
    ("Privacy: ", "no names, phone numbers or exact addresses are used in training or inference."),
    ("Fairness: ", "accuracy is monitored per language and per district to avoid systematically worse service for any group."),
])

# ----------------------------------------------------------------- 7 business model
page_break()
h("7. Business Model and Sustainability", 1)
para("Prajadhwani will always be free for citizens. Its customers are the institutions that must resolve grievances and plan "
     "local development; its value to them is faster redress, lower duplication, and evidence for budgeting.")
figure("f09_bmc", "Business Model Canvas")
h("7.1 Revenue model (indicative, for discussion)", 2)
table(["Stream", "Unit price (indicative)", "Illustrative annual potential"], [
    ("Urban local body licence", "₹12,000 per ward per year", "₹44 lakh for all 369 GBA wards"),
    ("District administration licence", "₹6 lakh per district per year", "₹1.86 crore for 31 districts"),
    ("Legislator dashboards", "₹60,000 per constituency per year", "₹1.5 crore for 224 assembly + 28 Lok Sabha seats"),
    ("Planning analytics for departments", "Custom", "Project-based"),
    ("Grants and CSR", "—", "Smart Cities Mission, Digital India, Google Cloud credits, CSR"),
], widths=[5.0, 5.2, 6.2])
caption("Table", "Indicative revenue streams (prices are proposals, not quotations)")
h("7.2 Cost structure and unit economics (indicative)", 2)
table(["Cost item", "Indicative cost", "Notes"], [
    ("AI processing per ticket (Gemini 2.5 Flash + translation)", "≈ ₹0.5 – 1.0", "Depends on text length and photo"),
    ("Gemma 4 router inference per ticket (self-hosted, batched)", "≈ ₹0.05 – 0.2", "Replaces repeated calls to a larger model"),
    ("Cloud infrastructure at 1 lakh tickets/month", "≈ ₹40,000 – 60,000 per month", "Firestore, Functions, Storage"),
    ("One-time router fine-tuning", "≈ ₹1 – 2 lakh", "Eligible for cloud research credits"),
    ("Support, onboarding and maintenance team", "≈ ₹6 lakh per month", "4–5 people at state scale"),
], widths=[7.2, 4.4, 4.8])
caption("Table", "Indicative operating costs")
para("At the scale of one city (≈1 lakh tickets per month) the platform costs about ₹90 lakh – 1 crore a year to run; licences "
     "from one city corporation group and a handful of districts cover this, making the model self-sustaining within about two years.")
h("7.3 Go-to-market", 2)
bullets([
    ("Phase 1 — Pilot (0–3 months): ", "one GBA ward with the ward office, a resident welfare association and a college NSS unit."),
    ("Phase 2 — Corporation (3–8 months): ", "scale to one GBA corporation (≈70 wards) with the Gemma 4 router live."),
    ("Phase 3 — Districts (8–12 months): ", "rural taluks via District Collector offices; State-level MoU and integration with existing redress portals."),
])

# ----------------------------------------------------------------- 8 impact
h("8. Impact and Alignment with National Priorities", 1)
table(["Outcome", "Indicator", "Pilot target"], [
    ("Faster redress", "Median time to resolution", "30% lower than ward baseline"),
    ("Correct routing", "Tickets reaching the right office first time", "≥ 90%"),
    ("Inclusion", "Share of reports in voice or non-English languages", "≥ 40%"),
    ("Transparency", "Tickets with a visible status update within 72 hours", "≥ 80%"),
    ("Better planning", "Hotspots with an AI solution card reviewed by officials", "100% of hotspots ≥ 5 reports"),
], widths=[3.8, 7.6, 5.0])
caption("Table", "Impact indicators")
table(["Programme / goal", "How Prajadhwani contributes"], [
    ("Swachh Bharat Mission", "Sanitation, garbage and drainage reports routed and tracked to closure"),
    ("Smart Cities Mission", "Ward-level civic data and hotspot analytics for planning"),
    ("Digital India", "Inclusive, voice-first, multilingual digital public service"),
    ("CPGRAMS / good governance", "Accountability, time-bound redress, citizen feedback loop"),
    ("SDG 6, 9, 11, 16", "Clean water & sanitation; infrastructure; sustainable cities; effective, accountable institutions"),
], widths=[5.2, 11.2])
caption("Table", "Alignment with national programmes and the Sustainable Development Goals")

# ----------------------------------------------------------------- 9 TRL & roadmap
page_break()
h("9. Technology Readiness and Implementation Roadmap", 1)
para("The platform is at TRL 5: all components run on the production cloud and have been verified end-to-end, the Android app is "
     "distributed to testers, and the AI pipeline has processed hundreds of test tickets. Reaching TRL 6 requires a pilot with "
     "real citizens and a real officer.")
figure("f11_trl", "Technology Readiness Level of Prajadhwani", width=13)
h("9.1 Pilot plan (TRL 5 → TRL 6)", 2)
table(["Element", "Plan"], [
    ("Location", "One GBA ward in Bengaluru, selected with the ward office"),
    ("Participants", "50+ citizens (via RWA and college volunteers), 1 ward officer, the area MLA's office"),
    ("Duration", "3 weeks of live use + 1 week of evaluation"),
    ("Success criteria", "≥ 100 real reports; ≥ 80% acknowledged within 72 hours; officer and citizen satisfaction surveys"),
    ("Data collected", "Routing correctness, time to acknowledge / resolve, language mix — feeds the Gemma 4 training set"),
], widths=[3.8, 12.6])
caption("Table", "Pilot plan")
figure("f12_roadmap", "12-month implementation roadmap")
h("9.2 Risks and mitigation", 2)
table(["Risk", "Likelihood", "Mitigation"], [
    ("Low adoption by officials", "Medium", "Co-design in the pilot; dashboards save officials time; MLA/MP sponsorship"),
    ("Mis-routing by AI", "Medium", "Confidence threshold, District Collector fallback, human re-assignment, monthly retraining"),
    ("Misuse / fake reports", "Medium", "Username accounts with rate limits, duplicate clustering, officer flagging"),
    ("Privacy breach", "Low", "No sensitive identifiers collected; role-scoped rules; anonymised public data; security audit"),
    ("Boundary changes (e.g. GBA delimitation)", "High", "Boundary data is versioned and both BBMP and GBA wards are supported"),
    ("Cloud cost growth", "Low", "Small self-hosted Gemma model, batching, cloud credits"),
], widths=[4.4, 2.4, 9.6])
caption("Table", "Risk register")

# ----------------------------------------------------------------- 10 compliance
h("10. Legal, Privacy and Compliance", 1)
bullets([
    ("Digital Personal Data Protection Act, 2023: ", "minimal personal data, clear purpose, consent at sign-up, and the ability to delete one's account."),
    ("Data localisation: ", "processing and storage in the Mumbai cloud region."),
    ("Accessibility: ", "voice input and 10 languages for low-literacy users; large touch targets."),
    ("Open standards: ", "OpenStreetMap-based mapping, open-weight AI model, exportable data for government systems."),
])

# ----------------------------------------------------------------- 11 team
h("11. Team", 1)
table(["Name", "Role", "Responsibility"], [
    ("[Member 1]", "Team lead", "Product, government liaison"),
    ("[Member 2]", "Mobile & cloud", "Flutter app, Firebase, Cloud Functions"),
    ("[Member 3]", "AI / ML", "Gemma 4 fine-tuning, evaluation"),
    ("[Member 4]", "Research & outreach", "Pilot coordination, surveys, impact measurement"),
    ("[Faculty mentor]", "Mentor", "Guidance and institutional coordination"),
], widths=[4.4, 4.0, 8.0])
caption("Table", "Team composition (to be completed)")

# ----------------------------------------------------------------- annexures
page_break()
h("Annexure A — Current Prototype Inventory", 1)
table(["Component", "Status"], [
    ("Android app (Flutter), 10 languages", "Implemented; distributed via Firebase App Distribution"),
    ("Username + password sign-in with security-question recovery", "Implemented"),
    ("Unified voice + text + photo report screen", "Implemented"),
    ("Instant ticket receipt and live status tracker", "Implemented"),
    ("AI transcription, translation, captioning, classification", "Implemented (Gemini 2.5 Flash, Cloud Translation)"),
    ("Geo-resolution: GBA 369 + BBMP 198 wards, 224 assemblies, taluks, districts", "Implemented"),
    ("Issue clustering and 4-agent hotspot analysis", "Implemented"),
    ("Official dashboards (MP, MLA, ward officer, Tahsildar, district officer)", "Implemented"),
    ("AI briefings and downloadable PDF reports", "Implemented"),
    ("Admin console and official approval workflow", "Implemented"),
    ("Public anonymised dashboard and maps", "Implemented"),
    ("Gemma 4 responsible-officer router with DC fallback", "Proposed — Phase 2"),
    ("Email and push notifications to assigned officers", "Proposed — Phase 2"),
], widths=[10.4, 6.0])
h("Annexure B — Glossary", 1)
table(["Term", "Meaning"], [
    ("BBMP", "Bruhat Bengaluru Mahanagara Palike — Bengaluru's former municipal corporation (198 wards)"),
    ("GBA", "Greater Bengaluru Authority — constituted 2025, with five city corporations and 369 wards"),
    ("CPGRAMS", "Centralised Public Grievance Redress and Monitoring System, Government of India"),
    ("DC", "Deputy Commissioner / District Collector — head of a district administration"),
    ("Gemma 4", "Google's open-weight model family (2026); E4B is a small edge-optimised variant"),
    ("LoRA / QLoRA", "Parameter-efficient fine-tuning methods that adapt a model with small trainable matrices"),
    ("TRL", "Technology Readiness Level, a 1–9 scale of technology maturity"),
    ("PII", "Personally identifiable information"),
], widths=[3.2, 13.2])
h("Annexure C — References", 1)
refs = [
    "Census of India 2011 — Karnataka population and literacy tables, Office of the Registrar General & Census Commissioner, India.",
    "National Commission on Population (2020). Report of the Technical Group on Population Projections 2011–2036.",
    "Karnataka Directorate of Economics and Statistics (2025). Population estimate for the Greater Bengaluru Authority area (as reported by Deccan Herald).",
    "Government of Karnataka (19 Nov 2025). Final ward delimitation notification for the five city corporations under the Greater Bengaluru Governance Act, 2024 (as reported by Deccan Herald / The South First).",
    "OpenCity (2025). “Decoding Bengaluru's Civic Complaints: A Deep Dive into BBMP Grievances Data 2025.” opencity.in.",
    "Department of Administrative Reforms and Public Grievances. CPGRAMS monthly reports, September–November 2025. Press Information Bureau.",
    "Google AI for Developers (2026). Gemma releases and model overview — Gemma 4. ai.google.dev/gemma.",
    "NVIDIA NeMo AutoModel documentation — Gemma model coverage and LoRA fine-tuning guides.",
    "Ministry of Education (2026). Seva First Innovation Challenge 2026 — sevafirst.mic.gov.in.",
    "Digital Personal Data Protection Act, 2023, Government of India.",
]
for i, ref in enumerate(refs, 1):
    q = doc.add_paragraph()
    q.paragraph_format.space_after = Pt(3)
    a = q.add_run(f"[{i}] ")
    a.bold = True
    a.font.size = Pt(9.5)
    b = q.add_run(ref)
    b.font.size = Pt(9.5)

# Ask Word to refresh fields (TOC) when the file is opened.
settings = doc.settings.element
upd = OxmlElement("w:updateFields")
upd.set(qn("w:val"), "true")
settings.append(upd)

doc.core_properties.title = "Prajadhwani — Concept Document (SFIC 2026)"
doc.core_properties.subject = "AI-powered citizen voice and grievance routing platform"
doc.core_properties.keywords = "Prajadhwani, SFIC 2026, Gemma 4, grievance redress, Karnataka"
doc.save(OUT)
print("saved", OUT)
